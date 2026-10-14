#!/usr/bin/env python3
"""SVG -> PPTX：每页嵌入原生 SVG（asvg:svgBlip）+ PNG 回退。
PowerPoint 2016+/365 显示矢量，可右键「转换为形状」得到可编辑文字；
旧版 / Keynote / LibreOffice 显示 PNG 回退。思路参考 zhx47/ppt-agent (Apache-2.0) 的 SvgImagePart，
但改为标准 svgBlip 扩展写法。用法: svg2pptx.py out.pptx a.svg b.svg ...
"""
import sys, subprocess, tempfile, os, re, copy
from lxml import etree
from pptx import Presentation
from pptx.util import Emu
from pptx.opc.package import Part
from pptx.opc.packuri import PackURI
from pptx.opc.constants import RELATIONSHIP_TYPE as RT

ASVG = "http://schemas.microsoft.com/office/drawing/2016/SVG/main"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"

def viewbox(svg):
    m = re.search(r'viewBox="\s*[\d.\-]+[ ,]+[\d.\-]+[ ,]+([\d.]+)[ ,]+([\d.]+)', svg)
    return (float(m.group(1)), float(m.group(2))) if m else (1280.0, 720.0)

def main(out, files):
    prs = Presentation()
    w, h = viewbox(open(files[0], encoding="utf-8").read())
    prs.slide_width = Emu(12192000); prs.slide_height = Emu(int(12192000 * h / w))
    layout = prs.slide_layouts[6]
    for i, f in enumerate(files, 1):
        svg = open(f, "rb").read()
        with tempfile.TemporaryDirectory() as td:
            png = os.path.join(td, "p.png")
            subprocess.run(["rsvg-convert", "-w", "1920", "-o", png, f], check=True)
            slide = prs.slides.add_slide(layout)
            pic = slide.shapes.add_picture(png, 0, 0, prs.slide_width, prs.slide_height)
        part = Part(PackURI(f"/ppt/media/svg{i}.svg"), "image/svg+xml", slide.part.package, svg)
        rid = slide.part.relate_to(part, RT.IMAGE)
        blip = pic._element.find(f".//{{{A}}}blip")
        ext_lst = etree.SubElement(blip, f"{{{A}}}extLst")
        ext = etree.SubElement(ext_lst, f"{{{A}}}ext", uri="{96DAC541-7B7A-43D3-8B79-37D633B846F1}")
        sb = etree.SubElement(ext, f"{{{ASVG}}}svgBlip", nsmap={"asvg": ASVG})
        sb.set(f"{{{R}}}embed", rid)
        pic.name = os.path.splitext(os.path.basename(f))[0]
    prs.save(out)

if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2:])
